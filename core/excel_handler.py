"""
excel_handler.py — Core Excel Handler
=======================================
Reads test case data from Excel and writes results back after each run.
Used by every flow script across all modules.

Current scope : CP (Broker) module
Future scope  : Developer, Admin modules
"""

from openpyxl import load_workbook, Workbook
from openpyxl.styles import PatternFill, Font
from colorama import Fore, init
from datetime import datetime
import os

init(autoreset=True)

# ─────────────────────────────────────────────────────────
# COLOURS FOR RESULT CELLS
# ─────────────────────────────────────────────────────────
GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
RED_FILL   = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
GREEN_FONT = Font(color="276221", bold=True)
RED_FONT   = Font(color="9C0006", bold=True)


class ExcelHandler:
    """
    Handles reading test data from Excel and writing results back.

    Usage:
        excel = ExcelHandler("test_cases/direct_booking.xlsx")
        rows  = excel.read_test_cases()
        excel.write_result(row_number=2, status="PASS", notes="Booking completed")
    """

    def __init__(self, filepath):
        self.filepath = filepath
        self._validate_file()

    # ─────────────────────────────────────────────────────
    # SETUP
    # ─────────────────────────────────────────────────────

    def _validate_file(self):
        """Check file exists and is a valid .xlsx before doing anything."""
        if not os.path.exists(self.filepath):
            print(Fore.RED + f"[EXCEL ERROR] File not found → {self.filepath}")
            raise FileNotFoundError(f"Excel file not found: {self.filepath}")

        if not self.filepath.endswith(".xlsx"):
            print(Fore.RED + f"[EXCEL ERROR] File must be .xlsx → {self.filepath}")
            raise ValueError("Only .xlsx files are supported")

        print(Fore.CYAN + f"[EXCEL] Loaded → {self.filepath}")


    # ─────────────────────────────────────────────────────
    # READ
    # ─────────────────────────────────────────────────────

    def read_test_cases(self, sheet_name=None):
        """
        Reads all rows from the Excel sheet.
        Returns a list of dicts — one dict per row, keys = column headers.

        Usage:
            rows = excel.read_test_cases()
            for row in rows:
                print(row["mobile"], row["payment_plan"])
        """
        wb = load_workbook(self.filepath)
        ws = wb[sheet_name] if sheet_name else wb.active

        headers = [cell.value for cell in ws[1]]
        rows = []

        for row in ws.iter_rows(min_row=2, values_only=True):
            # Skip completely empty rows
            if all(cell is None for cell in row):
                continue
            row_dict = dict(zip(headers, row))
            rows.append(row_dict)

        print(Fore.CYAN + f"[EXCEL] Read {len(rows)} test case(s) from '{ws.title}'")
        return rows


    def get_test_case(self, row_number, sheet_name=None):
        """
        Returns a single row as a dict by row number (1 = first data row, not header).

        Usage:
            row = excel.get_test_case(1)
            print(row["mobile"])
        """
        rows = self.read_test_cases(sheet_name)

        if row_number < 1 or row_number > len(rows):
            print(Fore.RED + f"[EXCEL ERROR] Row {row_number} does not exist. Total rows: {len(rows)}")
            raise IndexError(f"Row {row_number} out of range")

        return rows[row_number - 1]


    # ─────────────────────────────────────────────────────
    # WRITE
    # ─────────────────────────────────────────────────────

    def write_result(self, row_number, status, notes="", sheet_name=None):
        """
        Writes PASS/FAIL result back to Excel for a specific row.
        Adds timestamp and optional notes.
        Colours cell green for PASS, red for FAIL.

        Usage:
            excel.write_result(row_number=1, status="PASS", notes="Booking completed")
            excel.write_result(row_number=2, status="FAIL", notes="OTP page did not load")
        """
        wb = load_workbook(self.filepath)
        ws = wb[sheet_name] if sheet_name else wb.active

        headers = [cell.value for cell in ws[1]]

        # Auto-create result columns if they don't exist
        if "result" not in headers:
            ws.cell(row=1, column=len(headers) + 1, value="result")
            headers.append("result")

        if "timestamp" not in headers:
            ws.cell(row=1, column=len(headers) + 1, value="timestamp")
            headers.append("timestamp")

        if "notes" not in headers:
            ws.cell(row=1, column=len(headers) + 1, value="notes")
            headers.append("notes")

        result_col    = headers.index("result") + 1
        timestamp_col = headers.index("timestamp") + 1
        notes_col     = headers.index("notes") + 1

        excel_row = row_number + 1  # +1 because row 1 is header

        # Write values
        result_cell = ws.cell(row=excel_row, column=result_col, value=status)
        ws.cell(row=excel_row, column=timestamp_col, value=datetime.now().strftime("%d-%m-%Y %H:%M:%S"))
        ws.cell(row=excel_row, column=notes_col, value=notes)

        # Colour the result cell
        if status == "PASS":
            result_cell.fill = GREEN_FILL
            result_cell.font = GREEN_FONT
            print(Fore.GREEN + f"[EXCEL] Row {row_number} → PASS ✓")
        else:
            result_cell.fill = RED_FILL
            result_cell.font = RED_FONT
            print(Fore.RED + f"[EXCEL] Row {row_number} → FAIL ✗  ({notes})")

        wb.save(self.filepath)
        print(Fore.CYAN + f"[EXCEL] Saved → {self.filepath}")
